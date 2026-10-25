#!/usr/bin/env python3
"""hangul.py — 한글 파일 도구 (Claude가 한글 .hwp / .hwpx 파일을 읽고, 고치고, 만드는 도구)

쓰는 법
  python3 hangul.py info    문서.hwp                     # 무엇이 들었나(문단·표·그림·빈칸 수)
  python3 hangul.py read    문서.hwp [--out 글.md]        # 글 전체를 마크다운으로 (표 포함)
  python3 hangul.py replace 문서.hwpx 찾을말 바꿀말 [찾을말2 바꿀말2 ...] [--out 새.hwpx]
  python3 hangul.py fill    양식.hwpx 값.json [--out 채운.hwpx]   # 양식 빈칸 채우기
  python3 hangul.py merge   양식.hwpx 명단.csv [--out 폴더]       # 한 사람에 한 장씩 여러 장 만들기
  python3 hangul.py new     원고.md [--out 문서.hwpx]      # 마크다운 원고로 새 한글 문서
  python3 hangul.py convert 문서.hwp 문서.hwpx           # 형식 바꾸기 (hwp <-> hwpx)
  python3 hangul.py check   문서.hwpx                    # 한글에서 열릴지 구조 검사
  python3 hangul.py preview 문서.hwpx [--out 폴더]        # 쪽 그림(PNG)으로 미리 보기

준비물: pip install --break-system-packages python-hwpx olefile
        (미리 보기 그림은 playwright 또는 wkhtmltoimage 가 있으면 PNG, 없으면 HTML)

python-hwpx(Apache-2.0, https://github.com/airmang/python-hwpx) 위에서 돈다.
"""
import argparse
import csv
import json
import os
import re
import shutil
import struct
import subprocess
import sys
import warnings
import zlib

warnings.filterwarnings("ignore")
import logging  # noqa: E402

logging.disable(logging.CRITICAL)  # 라이브러리 내부 오류 로그는 숨기고, 우리가 한국어로 알린다


def _need_lib():
    try:
        import hwpx  # noqa: F401
    except ImportError:
        print("python-hwpx 가 없습니다. 먼저 이것을 실행하세요:\n"
              "  pip install --break-system-packages python-hwpx olefile", file=sys.stderr)
        sys.exit(2)


def _open(path):
    _need_lib()
    from hwpx import HwpxDocument
    if not os.path.exists(path):
        raise SystemExit(f"파일이 없습니다: {path}")
    if _kind(path) not in ("hwp", "hwpx"):
        raise SystemExit(f"한글 파일(.hwp, .hwpx)이 아닙니다: {path}")
    with open(path, "rb") as fh:
        head = fh.read(32)
    if head.startswith(b"HWP Document File V3"):
        raise SystemExit("한글 97 이전(한글 3.0) 형식의 아주 옛날 파일이라 열 수 없습니다. "
                         "한글이나 한컴오피스 뷰어에서 열어 '다른 이름으로 저장 → HWP 또는 HWPX'로 바꾼 파일을 주세요.")
    try:
        return HwpxDocument.open(path)
    except Exception as e:
        if _kind(path) == "hwpx" and "invalid Char" in str(e):
            # XML에 들어가면 안 되는 보이지 않는 글자가 든 HWPX → 그 글자만 빼고 연다 (실제 기안문 서식에서 나옴, 10-25)
            return HwpxDocument.open(_clean_hwpx_bytes(path))
        raise


def _clean_hwpx_bytes(path):
    import io
    import zipfile
    bad = re.compile(rb"[\x00-\x08\x0b\x0c\x0e-\x1f]")
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.endswith(".xml") or info.filename.endswith(".hpf"):
                data = bad.sub(b"", data)
            dst.writestr(info, data)
    return buf.getvalue()


def _kind(path):
    ext = os.path.splitext(path)[1].lower()
    return {".hwp": "hwp", ".hwpx": "hwpx"}.get(ext, ext.lstrip(".") or "?")


def _default_out(src, suffix, ext=None):
    base, e = os.path.splitext(os.path.basename(src))
    return f"{base}_{suffix}{ext or e}"


def _conversion_notes(doc):
    rep = getattr(doc, "conversion_report", None)
    if not rep:
        return []
    notes = []
    try:
        items = rep if isinstance(rep, (list, tuple)) else (
            rep.get("warnings") if isinstance(rep, dict) else getattr(rep, "warnings", None)) or []
        for it in items:
            notes.append(str(it)[:200])
    except Exception:
        notes.append(str(rep)[:300])
    return notes


# ------------------------------------------------------------ 예비 읽기 (아주 큰 .hwp 등)
def _fallback_hwp_text(path):
    """python-hwpx 가 못 연 .hwp 에서 글자만 직접 뽑는다 (서식·표 모양 없음)."""
    import olefile
    if not olefile.isOleFile(path):
        raise SystemExit("한글 5.0(.hwp) 형식이 아닙니다.")
    ole = olefile.OleFileIO(path)
    header = ole.openstream("FileHeader").read()
    flags = struct.unpack("<I", header[36:40])[0]
    compressed, encrypted, distribute = flags & 1, flags & 2, flags & 4
    if encrypted:
        raise SystemExit("암호가 걸린 한글 파일이라 읽을 수 없습니다. 한글에서 암호를 풀고 다시 주세요.")
    if distribute:
        raise SystemExit("배포용(읽기 전용) 한글 파일이라 본문이 잠겨 있습니다. 원본 파일을 받아야 합니다.")
    secs = sorted([e for e in ole.listdir() if e[0] == "BodyText"], key=lambda e: int(e[1][7:]))
    out = []
    for e in secs:
        data = ole.openstream(e).read()
        if compressed:
            data = zlib.decompress(data, -15)
        i = 0
        while i + 4 <= len(data):
            h = struct.unpack("<I", data[i:i + 4])[0]
            tag, size = h & 0x3FF, (h >> 20) & 0xFFF
            i += 4
            if size == 0xFFF:
                size = struct.unpack("<I", data[i:i + 4])[0]
                i += 4
            rec = data[i:i + size]
            i += size
            if tag == 67:  # HWPTAG_PARA_TEXT
                out.append(_decode_para_text(rec))
    return "\n".join(t for t in out)


def _decode_para_text(rec):
    chars, j = [], 0
    n = len(rec) // 2
    while j < n:
        c = struct.unpack("<H", rec[j * 2:j * 2 + 2])[0]
        if c in (1, 2, 3, 11, 12, 14, 15, 16, 17, 18, 21, 22, 23):  # 확장 컨트롤 (8글자 차지)
            j += 8
            continue
        if c in (4, 5, 6, 7, 8, 19, 20):  # 인라인 컨트롤 (8글자)
            j += 8
            continue
        if c == 9:
            chars.append("\t"); j += 8; continue
        if c in (10, 13):
            j += 1; continue
        if c == 24 or c == 30 or c == 31:
            chars.append(" "); j += 1; continue
        if c < 32:
            j += 1; continue
        if 0xD800 <= c <= 0xDBFF and j + 1 < n:  # 두 칸짜리 글자(이모지 등)
            c2 = struct.unpack("<H", rec[(j + 1) * 2:(j + 1) * 2 + 2])[0]
            if 0xDC00 <= c2 <= 0xDFFF:
                chars.append(chr(0x10000 + ((c - 0xD800) << 10) + (c2 - 0xDC00)))
                j += 2
                continue
        if 0xD800 <= c <= 0xDFFF:  # 짝 없는 반쪽 글자는 버린다 (실제 관공서 문서에서 나옴, 10-25)
            j += 1
            continue
        chars.append(chr(c))
        j += 1
    return "".join(chars)


# ------------------------------------------------------------ info
def cmd_info(a):
    try:
        doc = _open(a.file)
    except SystemExit:
        raise
    except Exception as e:
        return _info_fallback(a, e)
    from hwpx import TextExtractor  # noqa: F401
    paras = doc.paragraphs
    tables = list(doc.tables.all) if hasattr(doc.tables, "all") else []
    try:
        images = doc.list_images()
    except Exception:
        images = []
    try:
        fields = list(doc.fields.all)
    except Exception:
        fields = []
    text = doc.export_text() if hasattr(doc, "export_text") else "\n".join(p.text for p in paras)
    holes = sorted(set(re.findall(r"\{\{\s*([^{}]+?)\s*\}\}|\$\{\s*([^{}]+?)\s*\}", text)))
    holes = sorted({x or y for x, y in holes})
    info = {
        "파일": os.path.basename(a.file),
        "형식": _kind(a.file),
        "크기(KB)": round(os.path.getsize(a.file) / 1024, 1),
        "구역 수": len(doc.sections),
        "본문 문단 수": len(paras),
        "표 수": len(tables),
        "그림 수": len(images),
        "누름틀 수": len(fields),
        "누름틀 이름": [getattr(f, "name", str(f)) for f in fields][:50],
        "빈칸({{이름}}) 목록": holes,
        "글자 수(빈칸 뺌)": len(re.sub(r"\s", "", text)),
        "옮기면서 빠진 것": _conversion_notes(doc),
    }
    try:
        tmap = doc.get_table_map()
        info["표 목록"] = [
            {"번호": t.get("table_index"), "행": t.get("rows"), "열": t.get("cols"),
             "표 앞 글": (t.get("preceding_paragraph_text") or "")[:40]}
            for t in (tmap.get("tables", []) if isinstance(tmap, dict) else [])
        ]
    except Exception:
        pass
    print(json.dumps(info, ensure_ascii=False, indent=2))


def _info_fallback(a, err):
    if _kind(a.file) != "hwp":
        raise SystemExit(f"열 수 없는 파일입니다: {err}")
    text = _fallback_hwp_text(a.file)
    print(json.dumps({"파일": os.path.basename(a.file), "형식": "hwp", "읽기 방식": "예비(글자만)",
                      "이유": str(err)[:200], "글자 수(빈칸 뺌)": len(re.sub(r"\s", "", text)),
                      "문단 수": text.count("\n") + 1}, ensure_ascii=False, indent=2))


# ------------------------------------------------------------ read
def cmd_read(a):
    try:
        doc = _open(a.file)
        how = "보통"
        if a.plain:
            md = doc.export_markdown()
        else:
            try:  # 각주·미주·글상자 글·링크·색까지
                md = doc.export_rich_markdown()
            except Exception:
                md = doc.export_markdown()
    except SystemExit:
        raise
    except Exception as e:
        if _kind(a.file) != "hwp":
            raise SystemExit(f"열 수 없는 파일입니다: {e}")
        md = _fallback_hwp_text(a.file)
        how = f"예비(글자만, 이유: {str(e)[:120]})"
    md = re.sub(r"\n{3,}", "\n\n", md).strip() + "\n"
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"[읽기 방식: {how}] 저장: {a.out} ({len(md)}자)")
    else:
        if how != "보통":
            print(f"[읽기 방식: {how}]")
        print(md)


# ------------------------------------------------------------ 저장 (형식 유지 + 검사)
def _save(doc, out, src=None):
    out_dir = os.path.dirname(os.path.abspath(out))
    os.makedirs(out_dir, exist_ok=True)
    try:
        try:
            doc.save_to_path(out)
        except TypeError as e:
            # XML 주석(<!-- -->)이 든 문서에서 검사 도구가 멈추는 문제 → 주석만 빼고 다시 저장 (대량 시험에서 찾음, 10-25)
            if not _strip_xml_comments(doc):
                raise
            doc.save_to_path(out)
    except Exception as e:
        msg = str(e)
        if "open-safety" in msg or "refuses to open" in msg:
            hint = ("이 문서에는 도구가 아직 제대로 옮기지 못하는 부분이 있어서, 한글에서 안 열리는 파일이 "
                    "될까 봐 저장을 멈췄습니다. 원본은 그대로입니다.")
            if src and _kind(src) == "hwp":
                hint += " 한글 프로그램에서 이 파일을 '다른 이름으로 저장 → 파일 형식: HWPX'로 바꾼 뒤 다시 주면 대부분 됩니다."
            raise SystemExit(json.dumps({"결과": "저장 안 함", "이유": hint, "자세한 원인": msg[:300]},
                                        ensure_ascii=False, indent=2))
        if "cannot express" in msg:
            hint = ("이 문서에는 .hwp 형식으로 아직 옮길 수 없는 부분(예: 깨진 OLE 개체)이 있어 저장을 멈췄습니다. 원본은 그대로입니다. "
                    "결과를 .hwpx 로 저장하면 대부분 됩니다(--out 이름을 .hwpx 로).")
            raise SystemExit(json.dumps({"결과": "저장 안 함", "이유": hint, "자세한 원인": msg[:300]},
                                        ensure_ascii=False, indent=2))
        raise
    notes = []
    if _kind(out) == "hwpx":
        ok, msgs = _check_file(out)
        if not ok:
            notes.append("구조 검사 경고: " + "; ".join(msgs[:5]))
    return notes


def _strip_xml_comments(doc):
    removed = 0
    roots = []
    for attr in ("headers", "_headers", "_sections"):
        for part in getattr(doc, attr, None) or []:
            el = getattr(part, "element", None)
            if el is not None:
                roots.append(el)
    try:
        roots += [s.element for s in doc.sections]
    except Exception:
        pass
    for root in roots:
        for node in list(root.iter()):
            if not isinstance(node.tag, str):
                parent = node.getparent()
                if parent is not None:
                    if node.tail:
                        prev = node.getprevious()
                        if prev is not None:
                            prev.tail = (prev.tail or "") + node.tail
                        else:
                            parent.text = (parent.text or "") + node.tail
                    parent.remove(node)
                    removed += 1
    return removed


def _check_file(path):
    from hwpx import validate_package, validate_editor_open_safety
    msgs, ok = [], True
    try:
        r = validate_package(path)
        if not getattr(r, "ok", True):
            ok = False
            msgs += [str(x)[:160] for x in (getattr(r, "issues", None) or getattr(r, "errors", None) or [r])][:5]
    except Exception as e:
        ok = False
        msgs.append(f"패키지 검사 실패: {e}")
    try:
        r = validate_editor_open_safety(path)
        if not getattr(r, "ok", True):
            ok = False
            msgs += [str(x)[:160] for x in (getattr(r, "issues", None) or getattr(r, "errors", None) or [r])][:5]
    except Exception as e:
        ok = False
        msgs.append(f"열림 검사 실패: {e}")
    return ok, msgs


# ------------------------------------------------------------ replace
def cmd_replace(a):
    pairs = a.pairs
    if a.map:
        with open(a.map, encoding="utf-8") as f:
            m = json.load(f)
        pairs = pairs + [x for kv in m.items() for x in kv]
    if len(pairs) < 2 or len(pairs) % 2:
        raise SystemExit("찾을말과 바꿀말을 짝으로 주세요. 예: replace 문서.hwpx 홍길동 김철수")
    doc = _open(a.file)
    report = []
    total = 0
    for i in range(0, len(pairs), 2):
        find, to = pairs[i], pairs[i + 1]
        n = _replace_everywhere(doc, find, to)
        report.append({"찾을말": find, "바꿀말": to, "바꾼 곳": n})
        total += n
    out = a.out or _default_out(a.file, "고침")
    notes = _save(doc, out, a.file)
    print(json.dumps({"저장": out, "모두 바꾼 곳": total, "자세히": report,
                      "주의": notes + _conversion_notes(doc)}, ensure_ascii=False, indent=2))


def _iter_all_paragraphs(doc):
    """본문 문단과 표 안(표 속 표까지)의 문단을 모두."""
    def walk(paragraphs, depth=0):
        for p in paragraphs:
            yield p
            if depth > 8:
                continue
            try:
                for t in p.tables:
                    for row in t.rows:
                        for c in row.cells:
                            yield from walk(c.paragraphs, depth + 1)
            except Exception:
                pass
    yield from walk(doc.paragraphs)


def _replace_everywhere(doc, find, to):
    """라이브러리 바꾸기(머리말·각주 등 포함) + 놓친 곳(표 속 표 등)을 한 번 더."""
    n = doc.text.replace(find, to, everywhere=True)
    per_to = to.count(find) if find else 0
    for p in _iter_all_paragraphs(doc):
        try:
            txt = p.text or ""
        except Exception:
            continue
        if find not in txt:
            continue
        left = txt.count(find) - per_to * (txt.count(to) if to else 0)
        if left <= 0 or (to and to in txt and per_to):
            continue  # 이미 바뀐 문단
        new_txt = txt.replace(find, to)
        try:
            p.set_text_preserving_runs(new_txt)
        except Exception:
            p.text = new_txt
        n += txt.count(find)
    return n


def _replace_across_runs(doc, find, to):
    n = 0
    for p in _iter_all_paragraphs(doc):
        try:
            txt = p.text
        except Exception:
            continue
        if find in txt:
            k = txt.count(find)
            try:
                p.set_text_preserving_runs(txt.replace(find, to))
            except Exception:
                p.text = txt.replace(find, to)
            n += k
    return n


# ------------------------------------------------------------ fill
def _fill_doc(doc, data):
    """값 dict 로 문서의 빈칸을 채운다. (채운 곳, 못 채운 곳, 남은 빈칸)"""
    done, failed = [], []
    field_names = set()
    try:
        for x in doc.list_form_fields():
            nm = x.get("name") if isinstance(x, dict) else getattr(x, "name", None)
            if nm:
                field_names.add(nm)
    except Exception:
        pass

    table_map = {k: str(v) for k, v in data.items() if _is_table_key(k)}
    rest = {k: v for k, v in data.items() if not _is_table_key(k)}

    # 1) 표 칸: "성명 > right", "상속인/성명", "성명#2" (우리가 만든 칸 찾기)
    if table_map:
        grids = _collect_grids(doc)
        for key, val in table_map.items():
            ok, info = _fill_table_cell(grids, key, val)
            if ok:
                done.append({"어디": key, "방식": "표 칸", "칸": info})
            else:
                failed.append({"어디": key, "이유": info})

    cell_fields = None
    for key, val in rest.items():
        val = "" if val is None else str(val)
        # 2) 누름틀
        if key in field_names:
            try:
                doc.fill_form_field(val, name=key)
                done.append({"어디": key, "방식": "누름틀"})
                continue
            except Exception as e:
                failed.append({"어디": key, "이유": f"누름틀 채우기 실패: {e}"})
                continue
        # 2-1) 셀 필드(표 칸에 붙인 이름)
        if cell_fields is None:
            cell_fields = _cell_fields(_collect_grids(doc))
        if key in cell_fields:
            for t, pos in cell_fields[key]:
                t.set_cell_text(pos.anchor[0], pos.anchor[1], val, logical=True)
            done.append({"어디": key, "방식": "셀 필드", "곳": len(cell_fields[key])})
            continue
        # 3) {{이름}} / ${이름} 빈칸, 또는 그대로 쓴 글자
        n = 0
        if key.startswith("{{") or key.startswith("${") or key.startswith("<<"):
            tokens = [key]
        else:
            tokens = ["{{" + key + "}}", "{{ " + key + " }}", "${" + key + "}", "<<" + key + ">>"]
        for tok in tokens:
            n += _replace_everywhere(doc, tok, val)
        if n:
            done.append({"어디": key, "방식": "빈칸 글자", "곳": n})
        else:
            failed.append({"어디": key, "이유": "문서에 {{" + key + "}} 빈칸도, 같은 이름의 누름틀도 없음"})
    try:
        left = sorted(set(re.findall(r"\{\{[^{}]+\}\}", _all_text(doc))))
    except Exception:
        left = []
    return done, failed, left


# ------------------------------------------------------------ 표 칸 찾기 (실제 관공서 양식으로 맞춤, 10-25)
def _is_table_key(k):
    return (">" in k or "/" in k or re.search(r"#\d+\s*$", k.split(">")[0].strip()) is not None) and not k.startswith("{{")


def _norm_label(t):
    return re.sub(r"[\s:：·ㆍ]", "", t or "")


def _cell_text(cell):
    try:
        return "".join(p.text or "" for p in cell.paragraphs)
    except Exception:
        return getattr(cell, "text", "") or ""


def _collect_grids(doc):
    """문서의 모든 표(셀 안 표 포함)를 (표, 격자) 목록으로."""
    out = []

    def walk_table(t):
        try:
            grid = t.get_cell_map()
        except Exception:
            return
        out.append((t, grid))
        seen = set()
        for row in grid:
            for pos in row:
                if pos.anchor in seen:
                    continue
                seen.add(pos.anchor)
                try:
                    for q in pos.cell.paragraphs:
                        for sub in q.tables:
                            walk_table(sub)
                except Exception:
                    pass

    for p in doc.paragraphs:
        try:
            for t in p.tables:
                walk_table(t)
        except Exception:
            pass
    return out


def _anchors(grid):
    seen, res = set(), []
    for row in grid:
        for pos in row:
            if pos.anchor not in seen:
                seen.add(pos.anchor)
                res.append(pos)
    return res


def _append_in_cell(cell, value):
    try:
        para = cell.paragraphs[-1]
        runs = [x for x in para.runs if (x.text or "").strip()]
        cid = runs[-1].char_pr_id_ref if runs else None
        para.add_run("  " + value, char_pr_id_ref=cid) if cid is not None else para.add_run("  " + value)
        try:  # 글을 바꿨으니 줄 배치 기록을 지운다 (안 지우면 한컴에서 글자가 겹칠 수 있음)
            for el in list(para.element):
                if el.tag.endswith("linesegarray"):
                    para.element.remove(el)
        except Exception:
            pass
        return True
    except Exception:
        return False


def _cell_fields(grids):
    out = {}
    for t, grid in grids:
        for pos in _anchors(grid):
            nm = getattr(pos.cell, "field_name", None)
            if nm:
                out.setdefault(nm, []).append((t, pos))
    return out


def _parse_table_key(key):
    parts = [x.strip() for x in key.split(">")]
    head, dirs = parts[0], [d.lower() for d in parts[1:] if d]
    nth = None
    m = re.search(r"#(\d+)\s*$", head)
    if m:
        nth = int(m.group(1))
        head = head[:m.start()].strip()
    group = None
    if "/" in head:
        group, head = [x.strip() for x in head.rsplit("/", 1)]
    return group, head, nth, dirs


def _fill_table_cell(grids, key, value):
    group, label, nth, dirs = _parse_table_key(key)
    want = _norm_label(label)
    if not want:
        return False, "이름표가 비어 있음"
    cands = []
    for ti, (t, grid) in enumerate(grids):
        for pos in _anchors(grid):
            txt = _norm_label(_cell_text(pos.cell))
            if txt == want:
                cands.append((0, ti, pos))
            elif txt.startswith(want + "(") or txt.startswith(want + "<"):
                cands.append((1, ti, pos))  # "성명(대표자)" 같은 칸
    exact = [c for c in cands if c[0] == 0]
    cands = exact or cands
    if group:
        g = _norm_label(group)
        keep = []
        for _, ti, pos in cands:
            grid = grids[ti][1]
            r0 = pos.anchor[0]
            found = False
            for gp in _anchors(grid):
                gr, gc = gp.anchor
                rs = gp.span[0] if gp.span else 1
                if gc < pos.anchor[1] and gr <= r0 < gr + rs and _norm_label(_cell_text(gp.cell)) == g:
                    found = True
                    break
            if found:
                keep.append((_, ti, pos))
        cands = keep
        if not cands:
            return False, f"'{group}' 묶음 안에서 '{label}' 칸을 못 찾음"
    if not cands:
        return False, "그 이름표를 표에서 못 찾음"
    if nth is not None:
        if nth < 1 or nth > len(cands):
            return False, f"'{label}' 칸은 {len(cands)}개뿐인데 #{nth}를 찾음"
        cands = [cands[nth - 1]]
    if len(cands) > 1:
        groups = []
        for _, ti, pos in cands:
            grid = grids[ti][1]
            left = [gp for gp in _anchors(grid) if gp.anchor[1] < pos.anchor[1]
                    and gp.anchor[0] <= pos.anchor[0] < gp.anchor[0] + (gp.span[0] if gp.span else 1)
                    and _norm_label(_cell_text(gp.cell))]
            groups.append(_norm_label(_cell_text(left[0].cell)) if left else "")
        if group or not all(groups) or len(set(groups)) < len(groups):
            # 묶음 이름이 없거나 겹치면 순서 번호로 고르게 한다
            base = f"{group}/{label}" if group else label
            hint = ", ".join(f"{base}#{i + 1}" + (f"({g})" if g else "") for i, g in enumerate(groups))
        else:
            hint = ", ".join(f"{g}/{label}" for g in groups)
        return False, f"같은 이름표가 {len(cands)}개 있음 → 이렇게 골라 주세요: {hint}"
    _, ti, pos = cands[0]
    table, grid = grids[ti]
    nrows, ncols = len(grid), len(grid[0]) if grid else 0
    r, c = pos.anchor
    span = pos.span or (1, 1)

    def step(r, c, span, d):
        if d == "right":
            nc = c + span[1]
            return (r, nc) if nc < ncols else None
        if d == "left":
            return (r, c - 1) if c > 0 else None
        if d == "down":
            nr = r + span[0]
            return (nr, c) if nr < nrows else None
        if d == "up":
            return (r - 1, c) if r > 0 else None
        return "bad"

    if dirs:
        for d in dirs:
            nxt = step(r, c, span, d)
            if nxt == "bad":
                return False, f"모르는 방향: {d} (right/left/down/up 중 하나)"
            if nxt is None:
                return False, f"'{label}' 칸에서 {d} 쪽으로 더 갈 칸이 없음"
            p2 = grid[nxt[0]][nxt[1]]
            r, c = p2.anchor
            span = p2.span or (1, 1)
    else:
        # 방향을 안 적으면: 오른쪽 빈칸 → 아래 빈칸 순서로
        target = None
        for d in ("right", "down"):
            nxt = step(r, c, span, d)
            if nxt and nxt != "bad":
                p2 = grid[nxt[0]][nxt[1]]
                if not _norm_label(_cell_text(p2.cell)):
                    target = p2
                    break
        if target is None:
            # 관공서 양식에 흔한 모양: 이름표가 넓은 칸 위쪽에 있고 그 칸 안에 적는다 → 이름표 뒤에 이어 쓴다
            if _append_in_cell(pos.cell, value):
                return True, f"{r + 1}행 {c + 1}열 (이름표 칸 안, 이름표 뒤)"
            return False, f"'{label}' 칸 오른쪽·아래에 빈칸이 없음 (> right 처럼 방향을 적어 주세요)"
        r, c = target.anchor
    target_cell = grid[r][c]
    existing = _norm_label(_cell_text(target_cell.cell))
    if existing and dirs == []:
        return False, "채울 칸이 비어 있지 않음"
    try:
        table.set_cell_text(r, c, value, logical=True, split_merged=False)
    except TypeError:
        table.set_cell_text(r, c, value)
    return True, f"{r + 1}행 {c + 1}열"


def _all_text(doc):
    return "\n".join(p.text or "" for p in _iter_all_paragraphs(doc))


def cmd_fill(a):
    with open(a.data, encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, dict):
        raise SystemExit("값.json 은 {\"이름\": \"값\"} 모양이어야 합니다.")
    doc = _open(a.file)
    done, failed, left = _fill_doc(doc, data)
    out = a.out or _default_out(a.file, "채움")
    notes = _save(doc, out, a.file)
    print(json.dumps({"저장": out, "채운 곳": done, "못 채운 곳": failed, "아직 남은 빈칸": left,
                      "주의": notes}, ensure_ascii=False, indent=2))


def _ko_reason(r):
    return {"label not found": "그 이름표를 표에서 못 찾음",
            "ambiguous label": "같은 이름표가 여러 개라 어느 칸인지 모름",
            "path must include at least one direction": "방향(> right 등)을 안 적음"}.get(r, r)


# ------------------------------------------------------------ merge (여러 장)
def cmd_merge(a):
    rows = _load_rows(a.data)
    if not rows:
        raise SystemExit("명단에 줄이 없습니다.")
    out = a.out or (os.path.splitext(os.path.basename(a.file))[0] + "_여러장")
    os.makedirs(out, exist_ok=True)
    ext = os.path.splitext(a.file)[1].lower()
    pattern = a.name or ("{index:03d}" + ext)
    made, problems = [], []
    for i, row in enumerate(rows, 1):
        row = {str(k).strip(): ("" if v is None else str(v)) for k, v in row.items() if k}
        try:
            name = pattern.format(index=i, **row)
        except (KeyError, IndexError, ValueError):
            name = f"{i:03d}{ext}"
        name = re.sub(r'[\\/:*?"<>|]+', "_", name)
        if not os.path.splitext(name)[1]:
            name += ext
        doc = _open(a.file)
        done, failed, left = _fill_doc(doc, row)
        path = os.path.join(out, name)
        _save(doc, path)
        made.append(name)
        if failed or left:
            problems.append({"파일": name, "못 채운 곳": [f["어디"] for f in failed], "남은 빈칸": left})
    print(json.dumps({"폴더": out, "만든 장수": len(made), "파일": made[:30],
                      "문제 있는 장": problems[:30]}, ensure_ascii=False, indent=2))


def _load_rows(path):
    low = path.lower()
    if low.endswith((".xlsx", ".xlsm")):
        return _load_xlsx_rows(path)
    if low.endswith(".xls"):
        raise SystemExit("옛 엑셀(.xls)은 못 읽습니다. 엑셀에서 '다른 이름으로 저장 → Excel 통합 문서(.xlsx)'로 바꿔 주세요.")
    if low.endswith(".json"):
        with open(path, encoding="utf-8") as f:
            d = json.load(f)
        return d if isinstance(d, list) else [d]
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return [dict(r) for r in csv.DictReader(f)]
        except UnicodeDecodeError:
            continue
    raise SystemExit("명단 파일 글자를 읽지 못했습니다 (UTF-8 또는 CP949 CSV, 또는 JSON).")


def _load_xlsx_rows(path, sheet=None):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("엑셀 명단을 읽으려면 openpyxl 이 필요합니다 (pip install openpyxl)")
    import datetime
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]

    def cell_text(v):
        if v is None:
            return ""
        if isinstance(v, datetime.datetime):
            return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
        if isinstance(v, datetime.date):
            return v.strftime("%Y-%m-%d")
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    header, rows = None, []
    for r in ws.iter_rows(values_only=True):
        vals = [cell_text(v) for v in r]
        if not any(vals):
            continue
        if header is None:
            header = vals
            continue
        rows.append({h: (vals[i] if i < len(vals) else "") for i, h in enumerate(header) if h})
    wb.close()
    if header is None:
        raise SystemExit("엑셀 첫 시트가 비어 있습니다.")
    # 계산식 칸인데 계산된 값이 저장 안 된 경우(엑셀에서 한 번도 안 연 파일) 알린다
    try:
        wf = openpyxl.load_workbook(path, data_only=False, read_only=True)
        wv = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sf = wf[sheet] if sheet else wf.worksheets[0]
        sv = wv[sheet] if sheet else wv.worksheets[0]
        empty_formula = []
        for rf, rv in zip(sf.iter_rows(), sv.iter_rows()):
            for cf, cv in zip(rf, rv):
                if isinstance(cf.value, str) and cf.value.startswith("=") and cv.value is None:
                    empty_formula.append(cf.coordinate)
        wf.close(); wv.close()
        if empty_formula:
            print(f"(주의: 계산식 칸 {', '.join(empty_formula[:10])} 의 값이 비어 있습니다. "
                  f"엑셀에서 파일을 한 번 열고 저장하면 값이 들어갑니다.)", file=sys.stderr)
    except Exception:
        pass
    return rows


# ------------------------------------------------------------ new (마크다운 → 한글)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__)")


def cmd_new(a):
    _need_lib()
    from hwpx import HwpxDocument
    with open(a.md, encoding="utf-8") as f:
        src = f.read()
    blocks = _parse_md(src)
    doc = HwpxDocument.new()
    st = doc.styles
    ids = {
        "title": st.ensure_run(bold=True, size=20),
        "h1": st.ensure_run(bold=True, size=16),
        "h2": st.ensure_run(bold=True, size=13),
        "h3": st.ensure_run(bold=True, size=11),
        "bold": st.ensure_run(bold=True),
        "italic": st.ensure_run(italic=True),
        "quote": st.ensure_run(italic=True, color="#555555"),
        "th": st.ensure_run(bold=True),
        "caption": st.ensure_run(size=9, color="#555555"),
    }
    img_notes = []
    if a.font:
        try:
            st.replace_font(a.font) if hasattr(st, "replace_font") else None
        except Exception:
            pass
    first = [True]
    lists = []  # (kind, [paragraph indexes])
    fmt_jobs = []  # (index, kwargs)

    def body_index(p):
        for i, q in enumerate(doc.paragraphs):
            if q.element is p.element:
                return i
        return len(doc.paragraphs) - 1

    def para(text, run_id=None, inline=True):
        if first[0]:
            p = doc.paragraphs[0]
            first[0] = False
            if text:
                _fill_runs(p, text, run_id, ids, inline, reuse=True)
            return p
        # 앞 문단(가운데 맞춘 그림 등)의 모양을 물려받지 않게 기본 문단 모양(0)으로 (한컴 뷰어로 확인, 10-25)
        p = doc.add_paragraph("", include_run=False, para_pr_id_ref="0", inherit_style=False)
        _fill_runs(p, text, run_id, ids, inline)
        return p

    cur_list = None
    for b in blocks:
        t = b["type"]
        if t != "li" and cur_list is not None:
            lists.append(cur_list)
            cur_list = None
        if t == "h":
            lv = b["level"]
            key = "title" if lv == 1 else ("h1" if lv == 2 else ("h2" if lv == 3 else "h3"))
            p = para(b["text"], ids[key], inline=False)
            i = body_index(p)
            kw = {"spacing_before_pt": 6 if lv > 1 else 0, "spacing_after_pt": 4, "keep_with_next": True}
            if lv == 1:
                kw["alignment"] = "CENTER"
                kw["spacing_after_pt"] = 12
            fmt_jobs.append((i, kw))
        elif t == "p":
            p = para(b["text"])
            if b.get("align"):
                fmt_jobs.append((body_index(p), {"alignment": b["align"]}))
        elif t == "quote":
            p = para(b["text"], ids["quote"], inline=False)
            fmt_jobs.append((body_index(p), {"indent_left_mm": 8}))
        elif t == "li":
            p = para(b["text"])
            i = body_index(p)
            if cur_list is None or cur_list[0] != b["kind"]:
                if cur_list is not None:
                    lists.append(cur_list)
                cur_list = (b["kind"], [])
            cur_list[1].append((i, b["level"]))
        elif t == "hr":
            p = para("")
            fmt_jobs.append((body_index(p), {"bottom_border": True}))
        elif t == "pagebreak":
            p = para("")
            fmt_jobs.append((body_index(p), {"page_break_before": True}))
        elif t == "image":
            if first[0]:
                first[0] = False  # 첫 빈 문단 뒤에 그림을 붙인다
            r = _add_image_block(doc, b, os.path.dirname(os.path.abspath(a.md)))
            if r:
                img_notes.append(r)
                p = para(f"[그림 없음: {b['src']}]")
                continue
            if b["alt"]:
                # 그림과 설명이 다른 쪽으로 갈라지지 않게 (한컴 뷰어로 확인, 10-25)
                fmt_jobs.append((len(doc.paragraphs) - 1, {"keep_with_next": True}))
                p = para(b["alt"], ids["caption"], inline=False)
                fmt_jobs.append((body_index(p), {"alignment": "CENTER", "spacing_after_pt": 6}))
        elif t == "table":
            if first[0]:
                first[0] = False  # 첫 빈 문단 뒤에 표를 붙인다
            rows = b["rows"]
            nr, nc = len(rows), max(len(r) for r in rows)
            tb = doc.add_table(nr, nc)
            for r, row in enumerate(rows):
                for c in range(nc):
                    cell = row[c] if c < len(row) else ""
                    tb.set_cell_text(r, c, _strip_inline(cell))
                if r == 0 and b.get("header", True):
                    for c in range(nc):
                        try:
                            tb.set_cell_shading(0, c, "#E7EEF7")
                        except Exception:
                            pass
                        try:
                            for q in tb.cell(0, c).paragraphs:
                                for run in q.runs:
                                    run.char_pr_id_ref = ids["th"]
                        except Exception:
                            pass
            para("")  # 표 뒤 한 줄
    if cur_list is not None:
        lists.append(cur_list)

    for i, kw in fmt_jobs:
        try:
            st.apply_paragraph_format(paragraph_index=i, **kw)
        except Exception as e:
            print(f"(문단 모양 건너뜀 {i}: {e})", file=sys.stderr)
    for kind, items in lists:
        for lv in sorted({lv for _, lv in items}):
            idxs = [i for i, l in items if l == lv]
            try:
                if kind == "bullet":
                    st.apply_list_format(paragraph_indexes=idxs, kind="bullet", level=lv,
                                         bullet_char="●" if lv == 1 else "○")
                else:
                    st.apply_list_format(paragraph_indexes=idxs, kind="number", level=lv)
                if lv > 1:  # 2단계부터는 안쪽으로 들여 쓴다 (폴라리스 뷰어로 확인, 10-24)
                    st.apply_paragraph_format(paragraph_indexes=idxs, indent_left_mm=7.0 * (lv - 1))
            except Exception as e:
                print(f"(목록 모양 건너뜀: {e})", file=sys.stderr)

    out = a.out or (os.path.splitext(os.path.basename(a.md))[0] + ".hwpx")
    notes = _save(doc, out)
    print(json.dumps({"저장": out, "형식": _kind(out), "블록 수": len(blocks),
                      "표 수": sum(1 for b in blocks if b["type"] == "table"),
                      "그림 수": sum(1 for b in blocks if b["type"] == "image") - len(img_notes),
                      "주의": notes + img_notes},
                     ensure_ascii=False, indent=2))


_TEXT_WIDTH_MM = 150.0  # A4 기본 여백에서 글이 들어가는 폭(약)


def _add_image_block(doc, b, base_dir):
    """그림 한 장을 가운데에 넣는다. 문제가 있으면 한국어 이유를 돌려준다."""
    src = b["src"]
    path = src if os.path.isabs(src) else os.path.join(base_dir, src)
    if not os.path.exists(path):
        return f"그림 파일이 없습니다: {src}"
    try:
        from PIL import Image
    except ImportError:
        return "그림을 넣으려면 pillow 가 필요합니다 (pip install pillow)"
    try:
        im = Image.open(path)
        im.load()
    except Exception as e:
        return f"그림 파일을 열 수 없습니다: {src} ({e})"
    fmt = (im.format or "").lower()
    data = open(path, "rb").read()
    if fmt in ("jpeg", "jpg"):
        ext = "jpg"
    elif fmt in ("png", "gif", "bmp"):
        ext = fmt
    else:  # webp, tiff 등은 png 로 바꿔 넣는다
        import io
        buf = io.BytesIO()
        (im.convert("RGBA") if im.mode in ("P", "LA") else im).save(buf, "PNG")
        data, ext = buf.getvalue(), "png"
    w_px, h_px = im.size
    dpi = im.info.get("dpi", (96, 96))[0] or 96
    natural_mm = w_px / float(dpi) * 25.4
    if b.get("width"):
        width_mm = _TEXT_WIDTH_MM * b["width"] / 100.0 if b.get("unit") == "%" else b["width"]
    else:
        width_mm = natural_mm
    width_mm = max(10.0, min(width_mm, _TEXT_WIDTH_MM))
    height_mm = width_mm * h_px / float(w_px)
    if height_mm > 220.0:  # 한 쪽을 넘지 않게
        height_mm = 220.0
        width_mm = height_mm * w_px / float(h_px)
    doc.add_picture(data, ext, width_mm=round(width_mm, 1), height_mm=round(height_mm, 1), align="CENTER")
    return None


def _strip_inline(s):
    return re.sub(r"\*\*([^*]+)\*\*|__([^_]+)__|\*([^*]+)\*", lambda m: m.group(1) or m.group(2) or m.group(3), s)


def _fill_runs(p, text, run_id, ids, inline, reuse=False):
    """글을 런으로 넣는다. **굵게**, *기울임* 을 따로 런으로."""
    pieces = [(text, run_id)]
    if inline and run_id is None:
        pieces = []
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") or part.startswith("__"):
                pieces.append((part[2:-2], ids["bold"]))
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                pieces.append((part[1:-1], ids["italic"]))
            else:
                pieces.append((part, None))
    if reuse:
        # 첫 문단(구역 설정이 든 문단)은 기존 런을 살려 글만 넣는다
        p.text = "".join(x for x, _ in pieces)
        if len(pieces) == 1 and pieces[0][1] is not None:
            try:
                for r in p.runs:
                    if r.text:
                        r.char_pr_id_ref = pieces[0][1]
            except Exception:
                pass
        elif len(pieces) > 1:
            try:
                p.clear_text()
            except Exception:
                p.text = ""
            for x, rid in pieces:
                p.add_run(x, char_pr_id_ref=rid) if rid else p.add_run(x)
        return
    if not pieces:
        p.add_run("")
    for x, rid in pieces:
        p.add_run(x, char_pr_id_ref=rid) if rid else p.add_run(x)


def _parse_md(src):
    lines = src.replace("\r\n", "\n").split("\n")
    blocks, i, buf = [], 0, []

    def flush():
        if buf:
            blocks.append({"type": "p", "text": " ".join(s.strip() for s in buf)})
            buf.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush(); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            flush(); blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2).strip()}); i += 1; continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush(); blocks.append({"type": "hr"}); i += 1; continue
        if s in ("<쪽나눔>", "<pagebreak>", "\\newpage"):
            flush(); blocks.append({"type": "pagebreak"}); i += 1; continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append({"type": "table", "rows": rows, "header": True})
            continue
        m = re.match(r"^(\s*)([-*+•]|\d+[.)])\s+(.*)$", ln)
        if m:
            flush()
            indent = len(m.group(1).replace("\t", "    "))
            kind = "number" if m.group(2)[0].isdigit() else "bullet"
            blocks.append({"type": "li", "kind": kind, "level": min(1 + indent // 2, 3), "text": m.group(3).strip()})
            i += 1; continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*(?:\{\s*(?:폭|width)\s*=\s*([\d.]+)\s*(mm|%)?\s*\})?$", s)
        if m:
            flush()
            blocks.append({"type": "image", "alt": m.group(1).strip(), "src": m.group(2).strip().strip('"<>'),
                           "width": float(m.group(3)) if m.group(3) else None, "unit": m.group(4) or "mm"})
            i += 1; continue
        if s.startswith(">"):
            flush(); blocks.append({"type": "quote", "text": s.lstrip("> ").strip()}); i += 1; continue
        m = re.match(r"^<(가운데|오른쪽)>\s*(.*)$", s)
        if m:
            flush(); blocks.append({"type": "p", "text": m.group(2), "align": "CENTER" if m.group(1) == "가운데" else "RIGHT"}); i += 1; continue
        buf.append(ln)
        i += 1
    flush()
    return blocks


# ------------------------------------------------------------ convert
def cmd_convert(a):
    doc = _open(a.file)
    out = a.out
    if not out:
        out = os.path.splitext(os.path.basename(a.file))[0] + (".hwpx" if _kind(a.file) == "hwp" else ".hwp")
    if _kind(out) not in ("hwp", "hwpx"):
        if _kind(out) in ("md", "txt"):
            with open(out, "w", encoding="utf-8") as f:
                f.write(doc.export_markdown() if _kind(out) == "md" else doc.export_text())
            print(json.dumps({"저장": out}, ensure_ascii=False))
            return
        if _kind(out) == "html":
            with open(out, "w", encoding="utf-8") as f:
                f.write(doc.export_html())
            print(json.dumps({"저장": out}, ensure_ascii=False))
            return
        raise SystemExit("바꿀 수 있는 형식: .hwp .hwpx .md .txt .html")
    notes = _save(doc, out, a.file)
    print(json.dumps({"저장": out, "형식": f"{_kind(a.file)} → {_kind(out)}",
                      "옮기면서 빠진 것": _conversion_notes(doc), "주의": notes}, ensure_ascii=False, indent=2))


# ------------------------------------------------------------ check
def cmd_check(a):
    if _kind(a.file) == "hwp":
        doc = _open(a.file)
        tmp = os.path.join(os.path.dirname(os.path.abspath(a.out or ".")) if a.out else "/tmp", "_check_tmp.hwpx")
        doc.save_to_path(tmp)
        ok, msgs = _check_file(tmp)
        os.remove(tmp)
        print(json.dumps({"파일": a.file, "결과": "좋음" if ok else "문제 있음", "내용": msgs,
                          "참고": ".hwp 는 hwpx 로 바꿔서 검사했습니다", "옮기면서 빠진 것": _conversion_notes(doc)},
                         ensure_ascii=False, indent=2))
        return
    ok, msgs = _check_file(a.file)
    print(json.dumps({"파일": a.file, "결과": "좋음" if ok else "문제 있음", "내용": msgs,
                      "뜻": "구조 검사만 했습니다. 실제 한글 프로그램에서 열어 본 것은 아닙니다."},
                     ensure_ascii=False, indent=2))


# ------------------------------------------------------------ preview
def cmd_preview(a):
    _need_lib()
    from hwpx.tools.layout_preview import render_layout_preview
    src = a.file
    tmp = None
    if _kind(src) == "hwp":
        doc = _open(src)
        tmp = os.path.splitext(os.path.basename(src))[0] + "_미리보기용.hwpx"
        doc.save_to_path(tmp)
        src = tmp
    out = a.out or (os.path.splitext(os.path.basename(a.file))[0] + "_미리보기")
    os.makedirs(out, exist_ok=True)
    pv = render_layout_preview(src, mode="pages")
    pages = pv.page_html_documents
    if callable(pages):
        pages = pages()
    pages = list(pages or []) or [pv.html]
    made = []
    for n, html in enumerate(pages[: a.pages], 1):
        hp = os.path.join(out, f"쪽{n:02d}.html")
        with open(hp, "w", encoding="utf-8") as f:
            f.write(html)
        png = os.path.join(out, f"쪽{n:02d}.png")
        if _html_to_png(hp, png):
            made.append(png)
        else:
            made.append(hp)
    if tmp and os.path.exists(tmp):
        os.remove(tmp)
    print(json.dumps({"폴더": out, "쪽 수": len(pages), "만든 파일": made,
                      "뜻": "한글 프로그램 화면과 똑같지는 않은 '비슷한 모양' 미리 보기입니다.",
                      "경고": [str(w)[:120] for w in (pv.warnings or [])][:5]}, ensure_ascii=False, indent=2))


def _html_to_png(html_path, png_path):
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch()
            pg = b.new_page(viewport={"width": 900, "height": 1250}, device_scale_factor=1)
            pg.goto("file://" + os.path.abspath(html_path))
            pg.screenshot(path=png_path, full_page=True)
            b.close()
        return True
    except Exception:
        pass
    if shutil.which("wkhtmltoimage"):
        r = subprocess.run(["wkhtmltoimage", "--quiet", "--width", "900", html_path, png_path],
                           capture_output=True, timeout=120)
        return r.returncode == 0 and os.path.exists(png_path)
    return False


# ------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description="한글(.hwp/.hwpx) 파일 도구")
    sub = ap.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("info"); s.add_argument("file"); s.set_defaults(fn=cmd_info)
    s = sub.add_parser("read"); s.add_argument("file"); s.add_argument("--out")
    s.add_argument("--plain", action="store_true", help="색·링크 표시 없이 맨 글과 표만"); s.set_defaults(fn=cmd_read)
    s = sub.add_parser("replace"); s.add_argument("file"); s.add_argument("pairs", nargs="*")
    s.add_argument("--map", help="{\"찾을말\": \"바꿀말\"} JSON 파일"); s.add_argument("--out"); s.set_defaults(fn=cmd_replace)
    s = sub.add_parser("fill"); s.add_argument("file"); s.add_argument("data"); s.add_argument("--out"); s.set_defaults(fn=cmd_fill)
    s = sub.add_parser("merge"); s.add_argument("file"); s.add_argument("data"); s.add_argument("--out")
    s.add_argument("--name", help="파일 이름 틀, 예: '{index:03d}_{이름}.hwpx'"); s.set_defaults(fn=cmd_merge)
    s = sub.add_parser("new"); s.add_argument("md"); s.add_argument("--out"); s.add_argument("--font"); s.set_defaults(fn=cmd_new)
    s = sub.add_parser("convert"); s.add_argument("file"); s.add_argument("out", nargs="?"); s.set_defaults(fn=cmd_convert)
    s = sub.add_parser("check"); s.add_argument("file"); s.add_argument("--out"); s.set_defaults(fn=cmd_check)
    s = sub.add_parser("preview"); s.add_argument("file"); s.add_argument("--out"); s.add_argument("--pages", type=int, default=10)
    s.set_defaults(fn=cmd_preview)

    a = ap.parse_args()
    try:
        a.fn(a)
    except SystemExit:
        raise
    except BrokenPipeError:
        sys.exit(0)
    except Exception as e:  # 낯선 오류도 한국어 한 줄로
        name = type(e).__name__
        if name == "Hwp5Error":
            why = "암호·배포용(읽기 전용)·보안 문서라 열 수 없습니다. 한글에서 풀어서 다시 저장한 파일이 필요합니다."
        elif "surrogates not allowed" in str(e):
            why = "이 문서에는 깨진 특수 글자가 있어 서식째로는 고칠 수 없습니다. 'read'로 글자만은 읽을 수 있습니다. 한글에서 다시 저장한 파일이면 될 수 있습니다."
        elif name in ("BadZipFile", "XMLSyntaxError"):
            why = "파일이 깨졌거나 한글 파일이 아닙니다. 암호가 걸린 HWPX일 수도 있습니다."
        else:
            why = "예상하지 못한 오류입니다."
        print(json.dumps({"결과": "실패", "이유": why, "자세한 원인": f"{name}: {str(e)[:300]}"},
                         ensure_ascii=False, indent=2), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
